# -*- coding: utf-8 -*-
"""Genera konecta_invitados.xlsx (plantilla vacía). Uso: python src/make_excel.py"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L

F, OR, DARK, YEL = "Arial", "FD6C44", "141311", "FFF0A3"
N = 200  # filas disponibles
wb = Workbook(); ws = wb.active; ws.title = "Invitados"
thin = Side(style="thin", color="E2DDD3"); B = Border(bottom=thin, right=thin)
cols = [("#", 5), ("Nombre", 24), ("Empresa / proyecto", 24), ("Perfil", 22), ("Por qué encaja", 34), ("Lo invita", 13),
        ("Evento", 18), ("Estado", 14), ("Teléfono", 15), ("Correo", 26), ("¿Pagó ticket S/350?", 12), ("Notas", 30)]
ws["A1"] = "Konecta · Lista de invitados y afiliados"; ws["A1"].font = Font(name=F, size=16, bold=True, color="FFFFFF")
ws["A2"] = "Llenen solo las celdas blancas. Las columnas Perfil, Lo invita, Evento, Estado y ¿Pagó? tienen lista desplegable. La fila 5 es un ejemplo: bórrenla al empezar."
ws["A2"].font = Font(name=F, size=10, color="BDB6AA")
for r in (1, 2, 3):
    for c in range(1, len(cols) + 1): ws.cell(r, c).fill = PatternFill("solid", fgColor=DARK)
for i, (h, w) in enumerate(cols, 1):
    c = ws.cell(4, i, h); c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=OR)
    c.alignment = Alignment(vertical="center", wrap_text=True); ws.column_dimensions[L(i)].width = w
ws.row_dimensions[4].height = 32
for r in range(5, 5 + N):
    ws.cell(r, 1, r - 4).font = Font(name=F, size=10, color="6F695F")
    for c in range(2, len(cols) + 1):
        cell = ws.cell(r, c); cell.font = Font(name=F, size=10); cell.border = B
        cell.alignment = Alignment(vertical="top", wrap_text=c in (5, 12))
ejemplo = ["[Ejemplo] Nombre Apellido", "Marca del emprendedor", "Emprendedor con trayectoria",
           "5 años emprendiendo, marca propia, algo innovador", "Pablo", "Lanzamiento", "Invitado",
           "999 999 999", "correo@ejemplo.com", "No", "Fila de ejemplo: borrar"]
for c, v in enumerate(ejemplo, 2): ws.cell(5, c, v).font = Font(name=F, size=10, italic=True, color="8A8378")

def lista(opciones, col):
    d = DataValidation(type="list", formula1='"' + ",".join(opciones) + '"', allow_blank=True)
    ws.add_data_validation(d); d.add(f"{col}5:{col}{4 + N}")
lista(["Emprendedor con trayectoria", "Emprendedor etapa temprana", "Inversionista / fondo", "Empresa / corporativo", "Aliado", "Otro"], "D")
lista(["Pablo", "Sebastián", "Emilio"], "F")
lista(["Cena Lala", "Lanzamiento", "Cena + Lanzamiento", "Afiliado"], "G")
lista(["Por invitar", "Invitado", "Confirmado", "Afiliado", "No va"], "H")
lista(["Sí", "No"], "K")
rng = f"$A$5:$L${4 + N}"
for val, color in (("Confirmado", "E3F1EC"), ("Afiliado", "CDEBDD"), ("No va", "EEEEEE")):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$H5="{val}"'], fill=PatternFill("solid", fgColor=color)))
ws.freeze_panes = "C5"; ws.auto_filter.ref = f"A4:L{4 + N}"

# ---- Resumen (fórmulas; la fila 5 de ejemplo no cuenta)
s = wb.create_sheet("Resumen")
last = 4 + N
R, E, K = f"Invitados!$F$6:$F${last}", f"Invitados!$H$6:$H${last}", f"Invitados!$K$6:$K${last}"
s["A1"] = "Konecta · Resumen"; s["A1"].font = Font(name=F, size=16, bold=True)
s["A2"] = "Se calcula solo desde la hoja Invitados. Editar únicamente las celdas amarillas (metas y precio)."
s["A2"].font = Font(name=F, size=10, color="6F695F")
for i, h in enumerate(["Socio", "Meta afiliados", "Invitados", "Confirmados", "Afiliados", "Avance vs meta"], 1):
    c = s.cell(4, i, h); c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=OR)
for k, socio in enumerate(["Pablo", "Sebastián", "Emilio"]):
    r = 5 + k
    s.cell(r, 1, socio); s.cell(r, 2, 20).fill = PatternFill("solid", fgColor=YEL)
    s.cell(r, 3, f'=COUNTIFS({R},A{r},{E},"<>No va")-COUNTIFS({R},A{r},{E},"")')
    s.cell(r, 4, f'=COUNTIFS({R},A{r},{E},"Confirmado")')
    s.cell(r, 5, f'=COUNTIFS({R},A{r},{E},"Afiliado")')
    s.cell(r, 6, f"=IF(B{r}=0,0,E{r}/B{r})").number_format = "0%"
s["A8"] = "Total"
for c in "BCDE": s[f"{c}8"] = f"=SUM({c}5:{c}7)"
s["F8"] = "=IF(B8=0,0,E8/B8)"; s["F8"].number_format = "0%"
s["A10"] = "Meta de afiliados al lun 5 oct"; s["B10"] = 10; s["B10"].fill = PatternFill("solid", fgColor=YEL)
s["C10"] = "=IF(B10=0,0,E8/B10)"; s["C10"].number_format = "0%"
s["A12"] = "Lanzamiento · mié 14 oct"
s["A13"] = "Precio del ticket (S/)"; s["B13"] = 350; s["B13"].fill = PatternFill("solid", fgColor=YEL)
s["A14"] = "Tickets pagados"; s["B14"] = f'=COUNTIFS({K},"Sí")'
s["A15"] = "Recaudado (S/)"; s["B15"] = "=B13*B14"; s["B15"].number_format = '"S/" #,##0'
s["A17"] = "Fuente: metas (20 por socio, 10 al 5 oct) y precio S/350 acordados por los 3 socios, plan del 22 sep 2026."
for row in s.iter_rows(min_row=5, max_row=17):
    for c in row: c.font = Font(name=F, bold=c.row in (8, 12))
s["A17"].font = Font(name=F, size=9, italic=True, color="6F695F")
for c, w in zip("ABCDEF", (32, 15, 12, 13, 11, 15)): s.column_dimensions[c].width = w
wb.calculation.fullCalcOnLoad = True
wb.save("konecta_invitados.xlsx"); print("ok konecta_invitados.xlsx")
